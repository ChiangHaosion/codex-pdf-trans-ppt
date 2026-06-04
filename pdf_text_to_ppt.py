#!/usr/bin/env python3
from __future__ import annotations

import argparse
import io
import math
import re
from dataclasses import dataclass
from pathlib import Path

import fitz
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.text import MSO_ANCHOR, MSO_AUTO_SIZE, PP_ALIGN
from pptx.oxml import parse_xml
from pptx.oxml.ns import nsdecls, qn
from pptx.util import Inches, Pt


EMU_PER_INCH = 914400
PDF_POINTS_PER_INCH = 72
CONTENT_MARGIN_X_PT = 56
CONTENT_MARGIN_Y_PT = 28
TEXT_INSET_X_PT = 8
TEXT_INSET_Y_PT = 3
CAPACITY_SAFETY = 0.98
IMAGE_AREA_HEIGHT_PT = 190
APP_VERSION = "0.4.0"


@dataclass(frozen=True)
class ConversionProfile:
    name: str
    answer_fallback_page: int | None
    answer_min_scan_page: int
    answer_markers: tuple[str, ...]
    answer_keyword: str
    answer_keyword_count: int
    table_keywords: tuple[str, ...]
    table_placeholder: str


WORKBOOK_TABLE_KEYWORDS = (
    "统计表",
    "调查表",
    "图表",
    "比例",
    "用途 所占比例",
    "阅读方式",
    "每天阅读用时",
    "调查年级及人数",
    "经常使用",
    "偶尔使用",
    "从不使用",
)

CONVERSION_PROFILES = {
    "workbook": ConversionProfile(
        name="workbook",
        answer_fallback_page=29,
        answer_min_scan_page=20,
        answer_markers=("巩固训练答案",),
        answer_keyword="答案",
        answer_keyword_count=2,
        table_keywords=WORKBOOK_TABLE_KEYWORDS,
        table_placeholder="【图表占位：请参照原 PDF 图表】",
    ),
    "generic": ConversionProfile(
        name="generic",
        answer_fallback_page=None,
        answer_min_scan_page=1,
        answer_markers=("答案", "参考答案"),
        answer_keyword="答案",
        answer_keyword_count=3,
        table_keywords=("统计表", "调查表", "图表", "表格"),
        table_placeholder="【图表占位：请参照原 PDF 图表】",
    ),
}

TOPIC_TITLE_RE = re.compile(r"(专题[一二三四五六七八九十]+--([^\n（]+)（(\d+)）)")
ANSWER_CATEGORIES = {"说明文", "记叙文", "非连续性文本"}
NARRATIVE_TOPIC_ANSWER_KEYS = {
    3: "比喻",
    4: "拟人",
    5: "夸张",
    6: "对比",
    7: "反复",
    8: "排比",
    9: "反问",
    10: "动作描写",
    11: "语言描写",
    12: "神态描写",
    13: "心理描写",
    14: "外貌描写",
    15: "环境描写",
    16: "品析人物题",
    17: "词语含义题",
    18: "文段作用题",
    19: "文章主旨题",
}
ANSWER_SUBHEADINGS = {
    "比喻",
    "拟人",
    "夸张",
    "对比",
    "反复",
    "排比",
    "反问",
    "动作描写",
    "语言描写",
    "神态描写",
    "心理描写",
    "外貌描写",
    "环境描写",
}
TOPIC_BOILERPLATE_LINES = {
    "阅读答题技巧小纸条",
    "现代文阅读",
    "非连续性文本阅读",
    "说明文",
    "记叙文",
    "熟记技巧",
    "认真审题",
    "规范作答",
    "你本来就很棒",
    "你本来就很棒！",
}
INLINE_ANSWER_RED = 0xC62828
INLINE_TEXT_COLOR = 0x17212B
INLINE_TITLE_BLUE = 0x174EB8
INLINE_ANSWER_LEAD_RATIO = 0.52
INLINE_MIN_STANDALONE_CHARS = 90
ANSWER_ANIMATION_DURATION_MS = 500


@dataclass(frozen=True)
class ConversionSummary:
    pdf_path: Path
    pptx_path: Path
    source_pages: int
    processed_pages: int
    skipped_from_page: int | None
    slides: int
    layout: str
    answer_mode: str = "skip"
    matched_answers: int = 0
    unmatched_topics: int = 0
    unused_answers: int = 0


@dataclass(frozen=True)
class TopicUnit:
    title: str
    category: str
    number: int
    heading: str
    text: str


def pt_to_emu(value: float) -> int:
    return int(value / PDF_POINTS_PER_INCH * EMU_PER_INCH)


def int_to_rgb(value: int) -> RGBColor:
    return RGBColor((value >> 16) & 255, (value >> 8) & 255, value & 255)


def resolve_profile(profile: str | ConversionProfile) -> ConversionProfile:
    if isinstance(profile, ConversionProfile):
        return profile
    try:
        return CONVERSION_PROFILES[profile]
    except KeyError as error:
        choices = ", ".join(sorted(CONVERSION_PROFILES))
        raise ValueError(f"Unsupported profile: {profile}. Choose from: {choices}") from error


def span_is_bold(span: dict) -> bool:
    font = span.get("font", "").lower()
    return "bold" in font or "black" in font or "heavy" in font


def span_is_italic(span: dict) -> bool:
    font = span.get("font", "").lower()
    return "italic" in font or "oblique" in font


def clean_font_name(font_name: str) -> str:
    for marker in ("-Bold", "-Italic", "-Oblique", ",Bold", ",Italic"):
        font_name = font_name.replace(marker, "")
    return font_name.strip() or "Arial"


def add_text_block(
    slide,
    block: dict,
    scale: float,
    offset_x: float = 0,
    offset_y: float = 0,
    margin_left: float = 0,
    margin_top: float = 0,
) -> None:
    x0, y0, x1, y1 = block["bbox"]
    left = pt_to_emu((x0 - offset_x) * scale) + pt_to_emu(margin_left)
    top = pt_to_emu((y0 - offset_y) * scale) + pt_to_emu(margin_top)
    width = max(pt_to_emu((x1 - x0) * scale), pt_to_emu(6))
    height = max(pt_to_emu((y1 - y0) * scale * 1.08), pt_to_emu(6))

    box = slide.shapes.add_textbox(left, top, width, height)
    frame = box.text_frame
    frame.clear()
    frame.margin_left = 0
    frame.margin_right = 0
    frame.margin_top = 0
    frame.margin_bottom = 0
    frame.word_wrap = False
    frame.auto_size = MSO_AUTO_SIZE.TEXT_TO_FIT_SHAPE
    frame.vertical_anchor = MSO_ANCHOR.TOP

    first_paragraph = True
    for line in block.get("lines", []):
        paragraph = frame.paragraphs[0] if first_paragraph else frame.add_paragraph()
        first_paragraph = False
        paragraph.alignment = PP_ALIGN.LEFT
        paragraph.space_before = Pt(0)
        paragraph.space_after = Pt(0)
        paragraph.line_spacing = 1.0

        for span in line.get("spans", []):
            text = span.get("text", "")
            if not text:
                continue
            run = paragraph.add_run()
            run.text = text
            font = run.font
            font.size = Pt(max(span.get("size", 8) * scale, 3))
            font.name = clean_font_name(span.get("font", "Arial"))
            font.bold = span_is_bold(span)
            font.italic = span_is_italic(span)
            font.color.rgb = int_to_rgb(span.get("color", 0))


def has_text(block: dict) -> bool:
    return any(span.get("text", "").strip() for line in block.get("lines", []) for span in line.get("spans", []))


def split_blocks_by_vertical_gaps(blocks: list[dict], min_gap: float = 60) -> list[list[dict]]:
    if not blocks:
        return []

    sorted_blocks = sorted(blocks, key=lambda block: (block["bbox"][1], block["bbox"][0]))
    segments: list[list[dict]] = []
    current: list[dict] = []
    current_bottom: float | None = None

    for block in sorted_blocks:
        top = block["bbox"][1]
        bottom = block["bbox"][3]
        if current and current_bottom is not None and top - current_bottom >= min_gap:
            segments.append(current)
            current = []
        current.append(block)
        current_bottom = max(current_bottom or bottom, bottom)

    if current:
        segments.append(current)
    return segments


def segment_bbox(blocks: list[dict], page_rect: fitz.Rect) -> tuple[float, float, float, float]:
    x0 = min(block["bbox"][0] for block in blocks)
    y0 = min(block["bbox"][1] for block in blocks)
    x1 = max(block["bbox"][2] for block in blocks)
    y1 = max(block["bbox"][3] for block in blocks)

    pad_x = 4
    pad_y = 6
    return (
        max(0, x0 - pad_x),
        max(0, y0 - pad_y),
        min(page_rect.width, x1 + pad_x),
        min(page_rect.height, y1 + pad_y),
    )


def object_bbox(item: dict) -> tuple[float, float, float, float]:
    x0, y0, x1, y1 = item["bbox"]
    return float(x0), float(y0), float(x1), float(y1)


def combined_bbox(items: list[dict], page_rect: fitz.Rect) -> tuple[float, float, float, float]:
    if not items:
        return 0, 0, page_rect.width, page_rect.height

    x0 = min(object_bbox(item)[0] for item in items)
    y0 = min(object_bbox(item)[1] for item in items)
    x1 = max(object_bbox(item)[2] for item in items)
    y1 = max(object_bbox(item)[3] for item in items)
    pad_x = 4
    pad_y = 6
    return (
        max(0, x0 - pad_x),
        max(0, y0 - pad_y),
        min(page_rect.width, x1 + pad_x),
        min(page_rect.height, y1 + pad_y),
    )


def spec_bbox(segment: list[dict], images: list[dict], page_rect: fitz.Rect) -> tuple[float, float, float, float]:
    return combined_bbox([*segment, *images], page_rect)


def rect_overlap_area(first: tuple[float, float, float, float], second: tuple[float, float, float, float]) -> float:
    left = max(first[0], second[0])
    top = max(first[1], second[1])
    right = min(first[2], second[2])
    bottom = min(first[3], second[3])
    return max(0, right - left) * max(0, bottom - top)


def vertical_distance_to_box(center_y: float, box: tuple[float, float, float, float]) -> float:
    if box[1] <= center_y <= box[3]:
        return 0
    return min(abs(center_y - box[1]), abs(center_y - box[3]))


def assign_images_to_segments(image_blocks: list[dict], segments: list[list[dict]], page_rect: fitz.Rect) -> list[list[dict]]:
    assigned: list[list[dict]] = [[] for _ in segments]
    if not image_blocks or not segments:
        return assigned

    segment_boxes = [segment_bbox(segment, page_rect) for segment in segments]
    for image_block in image_blocks:
        image_box = object_bbox(image_block)
        image_center_y = (image_box[1] + image_box[3]) / 2
        best_index = 0
        best_score = -1.0
        for index, segment_box in enumerate(segment_boxes):
            overlap = rect_overlap_area(image_box, segment_box)
            if overlap > best_score:
                best_score = overlap
                best_index = index

        if best_score <= 0:
            best_index = min(
                range(len(segment_boxes)),
                key=lambda index: vertical_distance_to_box(image_center_y, segment_boxes[index]),
            )
        assigned[best_index].append(image_block)
    return assigned


def block_text(block: dict) -> str:
    if "_text" in block:
        return block["_text"]

    lines: list[str] = []
    for line in block.get("lines", []):
        text = "".join(span.get("text", "") for span in line.get("spans", [])).strip()
        if text:
            lines.append(text)
    return normalize_text(join_pdf_lines(lines))


def should_start_new_line(previous: str, current: str) -> bool:
    if not previous:
        return False
    if previous == "阅读答题技巧小纸条":
        return True
    if previous.startswith(("【", "一、", "二、", "三、")):
        return True
    if current.startswith(("【", "一、", "二、", "三、", "四、", "五、", "①", "②", "③", "④", "⑤", "⑥", "⑦", "⑧", "⑨", "⑩", "⑪", "⑫", "⑬", "任务")):
        return True
    if previous.endswith(("。", "？", "?", "！", "!", "：", ":")) and len(previous) < 42:
        return True
    return False


def join_pdf_lines(lines: list[str]) -> str:
    paragraphs: list[str] = []
    current = ""
    for raw_line in lines:
        line = raw_line.strip()
        if not line:
            continue
        if current and should_start_new_line(current, line):
            paragraphs.append(current)
            current = line
        else:
            current = f"{current}{line}" if current else line
    if current:
        paragraphs.append(current)
    return "\n".join(paragraphs).strip()


def normalize_text(text: str) -> str:
    text = re.sub(r"([（(])\s*\n\s*", r"\1", text)
    text = re.sub(r"\s*\n\s*([）)])", r"\1", text)
    text = re.sub(r"([\u4e00-\u9fff])\n([\u4e00-\u9fff])", r"\1\2", text)
    text = re.sub(r"([（(])\s+", r"\1", text)
    text = re.sub(r"\s+([）)])", r"\1", text)
    return text


def clean_extracted_lines(text: str) -> list[str]:
    lines: list[str] = []
    for raw_line in text.replace("\r", "\n").splitlines():
        line = raw_line.strip()
        if not line:
            continue
        line = re.sub(r"^[?？]\s*", "• ", line)
        line = re.sub(r"\s+", " ", line)
        lines.append(line)
    return lines


def clean_extracted_text(text: str) -> str:
    return normalize_text(join_pdf_lines(clean_extracted_lines(text)))


def clean_topic_text(text: str, title: str) -> str:
    lines: list[str] = []
    skipping_toc = False
    for line in clean_extracted_lines(text):
        stripped = line.strip()
        if not stripped:
            continue
        if skipping_toc:
            continue
        if stripped == title:
            continue
        if stripped in {"现代文阅读", "非连续性文本阅读"}:
            skipping_toc = True
            continue
        if stripped in TOPIC_BOILERPLATE_LINES:
            continue
        if all(piece in stripped for piece in ("熟记技巧", "认真审题", "规范作答")):
            continue
        if re.match(r"^\d+、(说明方法|常见题型|筛选主要信息题|材料信息探究题|提出建议题|图表类|材料类|图文结合类|标题含义题|文本概括题|赏析句子题|品析人物题|词语含义题|文段作用题|文章主旨题)", stripped):
            continue
        lines.append(stripped)
    return normalize_text(join_pdf_lines(lines))


def normalize_answer_heading(text: str) -> str:
    text = re.sub(r"^[•\-]\s*", "", text.strip())
    text = re.sub(r"^[①②③④⑤⑥⑦⑧⑨⑩、\s]+", "", text)
    text = re.sub(r"^\d+[、.．]\s*", "", text)
    text = text.strip("【】[] ")
    return text


def canonical_answer_key(text: str) -> str:
    text = normalize_answer_heading(text)
    text = text.replace("（小学考题不对此进行区分）", "")
    text = re.sub(r"[:：].*$", "", text)
    text = re.sub(r"[“”\"*？?]", "", text)
    text = re.sub(r"\s+", "", text)
    return text


def extract_topic_heading(text: str) -> str:
    bracketed = re.findall(r"【([^】]+)】", text)
    if not bracketed:
        return ""
    if len(bracketed) >= 2 and re.match(r"^\d", bracketed[0].strip()):
        return normalize_answer_heading(bracketed[1])
    return normalize_answer_heading(bracketed[0])


def split_topic_units(text: str) -> list[TopicUnit]:
    matches = list(TOPIC_TITLE_RE.finditer(text))
    units: list[TopicUnit] = []
    for index, match in enumerate(matches):
        start = match.start()
        end = matches[index + 1].start() if index + 1 < len(matches) else len(text)
        title = match.group(1)
        chunk = clean_topic_text(text[start:end], title)
        if not chunk:
            continue
        units.append(
            TopicUnit(
                title=title,
                category=match.group(2).strip(),
                number=int(match.group(3)),
                heading=extract_topic_heading(chunk),
                text=chunk,
            )
        )
    return units


def topic_answer_keys(unit: TopicUnit) -> list[str]:
    keys = [unit.title]
    if unit.category == "记叙文" and unit.number in NARRATIVE_TOPIC_ANSWER_KEYS:
        keys.append(f"{unit.category}:{canonical_answer_key(NARRATIVE_TOPIC_ANSWER_KEYS[unit.number])}")
    if unit.heading:
        keys.append(f"{unit.category}:{canonical_answer_key(unit.heading)}")
    keys.append(f"{unit.category}:{unit.number}")
    return list(dict.fromkeys(key for key in keys if key))


def answer_section_text(title: str, lines: list[str]) -> str:
    body = "\n".join(line for line in lines if line.strip()).strip()
    return f"{title}\n{body}".strip() if body else title.strip()


def add_answer_mapping(answer_map: dict[str, str], key: str, value: str) -> None:
    if key and value and key not in answer_map:
        answer_map[key] = value


def add_answer_section_mappings(answer_map: dict[str, str], category: str, number: int, title: str, lines: list[str]) -> None:
    section_text = answer_section_text(title, lines)
    add_answer_mapping(answer_map, f"{category}:{number}", section_text)
    add_answer_mapping(answer_map, f"{category}:{canonical_answer_key(title)}", section_text)

    current_subheading = ""
    current_lines: list[str] = []
    for line in lines:
        heading = normalize_answer_heading(line)
        if heading in ANSWER_SUBHEADINGS:
            if current_subheading and current_lines:
                add_answer_mapping(answer_map, f"{category}:{canonical_answer_key(current_subheading)}", answer_section_text(current_subheading, current_lines))
            current_subheading = heading
            current_lines = []
            continue
        if current_subheading:
            current_lines.append(line)
    if current_subheading and current_lines:
        add_answer_mapping(answer_map, f"{category}:{canonical_answer_key(current_subheading)}", answer_section_text(current_subheading, current_lines))


def parse_answer_units(text: str) -> dict[str, str]:
    answer_map: dict[str, str] = {}
    current_category = ""
    current_number: int | None = None
    current_title = ""
    current_lines: list[str] = []
    current_exact_title = ""
    current_exact_category = ""
    current_exact_number: int | None = None

    def flush_numbered_section() -> None:
        nonlocal current_number, current_title, current_lines
        if current_category and current_number is not None and current_title:
            add_answer_section_mappings(answer_map, current_category, current_number, current_title, current_lines)
        current_number = None
        current_title = ""
        current_lines = []

    def flush_exact_section() -> None:
        nonlocal current_exact_title, current_exact_category, current_exact_number, current_lines
        if current_exact_title and current_exact_number is not None:
            section_text = answer_section_text(current_exact_title, current_lines)
            add_answer_mapping(answer_map, current_exact_title, section_text)
            add_answer_mapping(answer_map, f"{current_exact_category}:{current_exact_number}", section_text)
        current_exact_title = ""
        current_exact_category = ""
        current_exact_number = None
        current_lines = []

    for line in clean_extracted_lines(text):
        topic_match = TOPIC_TITLE_RE.fullmatch(line)
        if topic_match:
            if current_exact_title:
                flush_exact_section()
            else:
                flush_numbered_section()
            current_exact_title = topic_match.group(1)
            current_exact_category = topic_match.group(2).strip()
            current_exact_number = int(topic_match.group(3))
            current_category = current_exact_category
            continue

        if current_exact_title:
            current_lines.append(line)
            continue

        if line in ANSWER_CATEGORIES:
            flush_numbered_section()
            current_category = line
            continue

        numbered_match = re.match(r"^(\d+)[.．]\s*(.+)$", line)
        if numbered_match and current_category:
            flush_numbered_section()
            current_number = int(numbered_match.group(1))
            current_title = normalize_answer_heading(numbered_match.group(2))
            current_lines = []
            continue

        if current_number is not None:
            current_lines.append(line)

    if current_exact_title:
        flush_exact_section()
    else:
        flush_numbered_section()
    return answer_map


def split_answer_parts(answer_text: str) -> tuple[str, str]:
    example_lines: list[str] = []
    training_lines: list[str] = []
    current: str | None = None

    for raw_line in answer_text.splitlines():
        line = raw_line.strip()
        if not line:
            continue
        marker_line = normalize_answer_heading(line)
        if marker_line in {"例题", "例题作答"} or marker_line.startswith("例题作答"):
            current = "example"
            line = re.sub(r"^[•\-]\s*", "", line)
            line = re.sub(r"^例题作答[:：]?", "", line).strip()
            if line:
                example_lines.append(line)
            continue
        if marker_line == "巩固训练" or "巩固训练" in marker_line and "答案" in marker_line:
            current = "training"
            line = re.sub(r"^[•\-]\s*", "", line)
            line = re.sub(r"^巩固训练\d*答案[:：]?", "", line).strip()
            if line:
                training_lines.append(line)
            continue
        if current == "example":
            example_lines.append(line)
        elif current == "training":
            training_lines.append(line)

    if not example_lines and not training_lines:
        training_lines = [line for line in answer_text.splitlines() if line.strip()]
    return "\n".join(example_lines).strip(), "\n".join(training_lines).strip()


def inline_block(
    text: str,
    font_size: int = 15,
    color: int = INLINE_TEXT_COLOR,
    bold: bool = False,
    role: str = "text",
) -> dict:
    return {
        "bbox": (0, 0, 720, 24),
        "lines": [],
        "_text": text.strip(),
        "_span": {
            "size": font_size,
            "font": "Arial",
            "color": color,
        },
        "_font_size": font_size,
        "_bold": bold,
        "_role": role,
    }


def split_topic_text_for_inline_answers(text: str) -> tuple[str, str]:
    marker_match = re.search(r"【\s*巩固训练\s*】", text)
    if not marker_match:
        return text.strip(), ""
    return text[: marker_match.start()].strip(), text[marker_match.start() :].strip()


def answer_for_topic(unit: TopicUnit, answer_map: dict[str, str], used_keys: set[str]) -> tuple[str, str | None]:
    for key in topic_answer_keys(unit):
        if key in answer_map:
            used_keys.add(key)
            return answer_map[key], key
    return "", None


def inline_answer_pair(label: str, answer_text: str) -> list[dict]:
    if not answer_text.strip():
        return []
    return [
        inline_block(label, font_size=15, color=INLINE_ANSWER_RED, bold=True, role="answer_label"),
        inline_block(answer_text, font_size=14, color=INLINE_ANSWER_RED, role="answer"),
    ]


def inline_groups_for_topic(unit: TopicUnit, answer_text: str) -> tuple[dict, list[list[dict]]]:
    example_question, training_question = split_topic_text_for_inline_answers(unit.text)
    example_answer, training_answer = split_answer_parts(answer_text)
    title = inline_block(unit.title, font_size=20, color=INLINE_TITLE_BLUE, bold=True, role="topic_title")
    groups: list[list[dict]] = []

    if example_question:
        group = [inline_block(example_question, font_size=15, color=INLINE_TEXT_COLOR, role="question")]
        group.extend(inline_answer_pair("例题参考答案", example_answer))
        if answer_text and not example_answer and not training_answer and not training_question:
            group.extend(inline_answer_pair("参考答案", answer_text))
        groups.append(group)

    if training_question:
        group = [inline_block(training_question, font_size=15, color=INLINE_TEXT_COLOR, role="question")]
        group.extend(inline_answer_pair("巩固训练参考答案", training_answer))
        if answer_text and not example_answer and not training_answer:
            group.extend(inline_answer_pair("参考答案", answer_text))
        groups.append(group)

    if training_answer and not training_question:
        if groups:
            groups[-1].extend(inline_answer_pair("巩固训练参考答案", training_answer))
        else:
            groups.append(inline_answer_pair("巩固训练参考答案", training_answer))

    if answer_text and not groups:
        groups.append(inline_answer_pair("参考答案", answer_text))

    return title, [[block for block in group if block_text(block)] for group in groups if any(block_text(block) for block in group)]


def inline_blocks_for_topic(unit: TopicUnit, answer_text: str) -> list[dict]:
    title, groups = inline_groups_for_topic(unit, answer_text)
    blocks: list[dict] = [title]
    for group in groups:
        blocks.extend(group)
    return blocks


def segment_is_title_only(segment: list[dict]) -> bool:
    text_blocks = [block for block in segment if block_text(block).strip()]
    return bool(text_blocks) and all(block.get("_role") == "topic_title" for block in text_blocks)


def split_first_chunk_for_prefix(
    chunks: list[dict],
    prefix: list[dict],
    width_pt: float,
    max_height_pt: float,
    base_size: int,
) -> list[dict]:
    if not prefix or not chunks or estimate_blocks_height([*prefix, chunks[0]], width_pt, base_size) <= max_height_pt:
        return chunks
    available_height = max_height_pt - estimate_blocks_height(prefix, width_pt, base_size)
    if available_height <= 0:
        return chunks
    return [*split_block_by_capacity(chunks[0], width_pt, available_height, base_size), *chunks[1:]]


def compact_block_length(block: dict) -> int:
    return len("".join(block_text(block).split()))


def rebalance_sparse_inline_chunks(
    chunks: list[dict],
    width_pt: float,
    max_height_pt: float,
    base_size: int,
    start_index: int = 0,
) -> list[dict]:
    index = max(start_index, 0)
    balanced = list(chunks)
    while index + 1 < len(balanced):
        if compact_block_length(balanced[index]) >= INLINE_MIN_STANDALONE_CHARS:
            index += 1
            continue

        combined = make_text_block_like(
            balanced[index],
            f"{block_text(balanced[index])}\n{block_text(balanced[index + 1])}",
        )
        replacement = split_block_by_capacity(combined, width_pt, max_height_pt, base_size)
        if not replacement or block_text(replacement[0]) == block_text(balanced[index]):
            index += 1
            continue
        balanced = [*balanced[:index], *replacement, *balanced[index + 2 :]]
    return balanced


def split_answer_blocks_for_inline(
    answer_blocks: list[dict],
    width_pt: float,
    max_height_pt: float,
    base_size: int,
) -> tuple[list[dict], list[list[dict]]]:
    if not answer_blocks:
        return [], []
    label = answer_blocks[0]
    body_blocks = answer_blocks[1:]
    if not body_blocks:
        return [label], []

    label_height = estimate_blocks_height([label], width_pt, base_size)
    lead_body_height = max(max_height_pt * INLINE_ANSWER_LEAD_RATIO - label_height, 55)
    first_body_chunks = split_block_by_capacity(body_blocks[0], width_pt, lead_body_height, base_size)
    lead = [label]
    rest_blocks: list[dict] = []
    if first_body_chunks:
        lead.append(first_body_chunks[0])
        rest_blocks.extend(first_body_chunks[1:])
    rest_blocks.extend(body_blocks[1:])
    rest_segments = (
        split_segments_by_capacity([rest_blocks], width_pt=width_pt, max_height_pt=max_height_pt, base_size=base_size)
        if rest_blocks
        else []
    )
    return lead, rest_segments


def split_inline_plain_group_by_capacity(
    group: list[dict],
    prefix: list[dict],
    width_pt: float,
    max_height_pt: float,
    base_size: int,
) -> list[list[dict]]:
    segments: list[list[dict]] = []
    current = list(prefix)
    for block in group:
        chunks = split_block_by_capacity(block, width_pt, max_height_pt, base_size)
        if segment_is_title_only(current):
            chunks = split_first_chunk_for_prefix(chunks, current, width_pt, max_height_pt, base_size)
        for chunk in chunks:
            if current and estimate_blocks_height([*current, chunk], width_pt, base_size) > max_height_pt:
                if segment_is_title_only(current):
                    current.append(chunk)
                    segments.append(current)
                else:
                    segments.append(current)
                    current = [chunk]
            else:
                current.append(chunk)

    if current:
        segments.append(current)
    return [segment for segment in segments if not segment_is_title_only(segment)]


def split_inline_answer_group_by_capacity(
    group: list[dict],
    prefix: list[dict],
    width_pt: float,
    max_height_pt: float,
    base_size: int,
) -> list[list[dict]]:
    answer_index = next(
        (index for index, block in enumerate(group) if block.get("_role") in {"answer_label", "answer"}),
        len(group),
    )
    question_blocks = group[:answer_index]
    answer_blocks = group[answer_index:]
    if not answer_blocks:
        return split_inline_plain_group_by_capacity(group, prefix, width_pt, max_height_pt, base_size)

    question_block = question_blocks[-1] if question_blocks else None
    leading_question_blocks = question_blocks[:-1]
    segments: list[list[dict]] = []
    active_prefix = list(prefix)

    if leading_question_blocks:
        leading_segments = split_inline_plain_group_by_capacity(
            leading_question_blocks,
            active_prefix,
            width_pt,
            max_height_pt,
            base_size,
        )
        if leading_segments:
            segments.extend(leading_segments[:-1])
            active_prefix = leading_segments[-1]

    answer_lead, answer_rest_segments = split_answer_blocks_for_inline(answer_blocks, width_pt, max_height_pt, base_size)
    if question_block is None:
        first_segment = [*active_prefix, *answer_lead]
        if estimate_blocks_height(first_segment, width_pt, base_size) > max_height_pt and active_prefix:
            segments.append(active_prefix)
            first_segment = list(answer_lead)
        segments.append(first_segment)
    else:
        question_chunks = split_block_by_capacity(question_block, width_pt, max_height_pt, base_size)
        if active_prefix:
            question_chunks = split_first_chunk_for_prefix(question_chunks, active_prefix, width_pt, max_height_pt, base_size)
        question_chunks = rebalance_sparse_inline_chunks(
            question_chunks,
            width_pt,
            max_height_pt,
            base_size,
            start_index=1 if active_prefix else 0,
        )

        answer_lead_height = estimate_blocks_height(answer_lead, width_pt, base_size)
        while question_chunks:
            final_prefix = active_prefix if len(question_chunks) == 1 else []
            final_height = estimate_blocks_height([*final_prefix, question_chunks[-1], *answer_lead], width_pt, base_size)
            if final_height <= max_height_pt:
                break
            reserved_height = estimate_blocks_height([*final_prefix, *answer_lead], width_pt, base_size)
            final_capacity = max(max_height_pt - reserved_height, block_font_size(question_chunks[-1], base_size) * 2.2)
            replacement_chunks = split_block_by_capacity(question_chunks[-1], width_pt, final_capacity, base_size)
            if len(replacement_chunks) <= 1:
                break
            question_chunks = [*question_chunks[:-1], *replacement_chunks]
            if answer_lead_height >= max_height_pt:
                break

        if question_chunks:
            for index, chunk in enumerate(question_chunks):
                segment: list[dict] = []
                if index == 0:
                    segment.extend(active_prefix)
                segment.append(chunk)
                if index == len(question_chunks) - 1:
                    segment.extend(answer_lead)
                segments.append(segment)
        elif answer_lead:
            segments.append([*active_prefix, *answer_lead])

    for rest_segment in answer_rest_segments:
        if segments and estimate_blocks_height([*segments[-1], *rest_segment], width_pt, base_size) <= max_height_pt:
            segments[-1].extend(rest_segment)
        else:
            segments.append(rest_segment)

    return [segment for segment in segments if segment and not segment_is_title_only(segment)]


def split_inline_topic_by_capacity(
    title: dict,
    groups: list[list[dict]],
    width_pt: float,
    max_height_pt: float,
    base_size: int = 15,
) -> list[list[dict]]:
    segments: list[list[dict]] = []
    current: list[dict] = [title]

    for group in groups:
        if not group:
            continue
        if estimate_blocks_height([*current, *group], width_pt, base_size) <= max_height_pt:
            current.extend(group)
            continue

        if current and not segment_is_title_only(current):
            segments.append(current)
            current = []

        if estimate_blocks_height([*current, *group], width_pt, base_size) <= max_height_pt:
            current.extend(group)
            continue

        split_segments = split_inline_answer_group_by_capacity(group, current, width_pt, max_height_pt, base_size)
        if split_segments:
            segments.extend(split_segments[:-1])
            current = split_segments[-1]
        else:
            current = []

    if current and not segment_is_title_only(current):
        segments.append(current)
    return segments


def build_inline_answer_slide_specs(
    doc: fitz.Document,
    answer_start_page: int,
    slide_width_pt: float,
    slide_height_pt: float,
) -> tuple[list[tuple[list[dict], list[dict]]], int, int, int, int]:
    question_text = "\n".join(doc[index].get_text() for index in range(0, answer_start_page - 1))
    answer_text = "\n".join(doc[index].get_text() for index in range(answer_start_page - 1, doc.page_count))
    topic_units = split_topic_units(question_text)
    answer_map = parse_answer_units(answer_text)
    used_answer_keys: set[str] = set()
    slide_specs: list[tuple[list[dict], list[dict]]] = []
    matched_answers = 0
    unmatched_topics = 0
    content_width_pt = slide_width_pt - CONTENT_MARGIN_X_PT * 2 - TEXT_INSET_X_PT * 2
    content_height_pt = (slide_height_pt - CONTENT_MARGIN_Y_PT * 2 - TEXT_INSET_Y_PT * 2) * CAPACITY_SAFETY

    for unit in topic_units:
        answer_text_for_unit, _key = answer_for_topic(unit, answer_map, used_answer_keys)
        if answer_text_for_unit:
            matched_answers += 1
        else:
            unmatched_topics += 1
        title, groups = inline_groups_for_topic(unit, answer_text_for_unit)
        segments = split_inline_topic_by_capacity(title, groups, width_pt=content_width_pt, max_height_pt=content_height_pt, base_size=15)
        for segment in segments:
            slide_specs.append((segment, []))

    unused_answers = len(set(answer_map) - used_answer_keys)
    return slide_specs, len(topic_units), matched_answers, unmatched_topics, unused_answers


def blocks_text_length(blocks: list[dict]) -> int:
    return sum(len(block_text(block)) for block in blocks)


def make_text_block_like(source: dict, text: str) -> dict:
    output = {
        "bbox": source["bbox"],
        "lines": source.get("lines", []),
        "_text": text,
        "_span": first_span(source) or {},
    }
    for key in ("_font_size", "_bold", "_role"):
        if key in source:
            output[key] = source[key]
    return output


def is_table_or_chart_text(text: str, profile: ConversionProfile) -> bool:
    compact = " ".join(text.split())
    percent_count = compact.count("%")
    digit_count = sum(char.isdigit() for char in compact)
    return percent_count >= 3 or digit_count >= 10 or any(keyword in compact for keyword in profile.table_keywords)


def replace_table_text_with_placeholders(blocks: list[dict], profile: ConversionProfile) -> list[dict]:
    output: list[dict] = []
    table_mode = 0
    for block in blocks:
        text = block_text(block)
        if table_mode and "阅读答题技巧小纸条" in text:
            table_mode = 0
        if is_table_or_chart_text(text, profile):
            if not output or block_text(output[-1]) != profile.table_placeholder:
                output.append(make_text_block_like(block, profile.table_placeholder))
            table_mode = 8
            continue
        if table_mode:
            compact = " ".join(text.split())
            digit_count = sum(char.isdigit() for char in compact)
            looks_like_table_row = len(compact) < 90 or digit_count >= 4 or "%" in compact
            if looks_like_table_row:
                table_mode -= 1
                continue
            table_mode = 0
        output.append(block)
    return output


def should_merge_text_blocks(previous_text: str, current_text: str) -> bool:
    if not previous_text or not current_text:
        return False
    if current_text.startswith(("参考答案", "例题参考答案", "巩固训练参考答案")):
        return False
    if previous_text.startswith(("参考答案", "例题参考答案", "巩固训练参考答案")):
        return False
    previous_tail = previous_text[-1]
    current_head = current_text[0]
    if previous_tail in "。！？!?：:；;）)】》”’…":
        return False
    if current_head in "【一二三四五六七八九十①②③④⑤⑥⑦⑧⑨⑩⑪⑫⑬任务":
        return False
    return "\u4e00" <= previous_tail <= "\u9fff" and "\u4e00" <= current_head <= "\u9fff"


def merge_adjacent_text_blocks(blocks: list[dict]) -> list[dict]:
    merged: list[dict] = []
    for block in blocks:
        text = block_text(block)
        if merged and should_merge_text_blocks(block_text(merged[-1]), text):
            previous = merged[-1]
            previous_text = block_text(previous)
            merged[-1] = make_text_block_like(previous, f"{previous_text}{text}")
        else:
            merged.append(block)
    return merged


def visual_units(text: str) -> float:
    units = 0.0
    for char in text:
        if char.isspace():
            units += 0.25
        elif ord(char) < 128:
            units += 0.55
        else:
            units += 1.0
    return units


def estimate_wrapped_lines(text: str, font_size: float, width_pt: float) -> int:
    chars_per_line = max(width_pt / (font_size * 0.92), 8)
    lines = 0
    for raw_line in text.splitlines() or [text]:
        line = raw_line.strip()
        if not line:
            continue
        lines += max(1, math.ceil(visual_units(line) / chars_per_line))
    return max(lines, 1)


def block_font_size(block: dict, base_size: int) -> int:
    if "_font_size" in block:
        return block["_font_size"]
    text = block_text(block)
    span = first_span(block) or {}
    source_size = span.get("size", 10)
    is_heading = source_size >= 13 or span_is_bold(span)
    return base_size + 2 if is_heading and len(text) < 90 else base_size


def estimate_blocks_height(blocks: list[dict], width_pt: float, base_size: int = 18) -> float:
    height = 0.0
    for block in blocks:
        text = block_text(block)
        if not text:
            continue
        font_size = block_font_size(block, base_size)
        height += estimate_wrapped_lines(text, font_size, width_pt) * font_size * 1.28
        height += 7
    return height


def split_text_by_height(text: str, width_pt: float, max_height_pt: float, font_size: int) -> list[str]:
    chunks: list[str] = []
    current_lines: list[str] = []

    def current_height(lines: list[str]) -> float:
        return estimate_wrapped_lines("\n".join(lines), font_size, width_pt) * font_size * 1.28 + 7

    for raw_line in split_text_into_sentences(text):
        line = raw_line.strip()
        if not line:
            continue
        pieces: list[str] = []
        chars_per_piece = max(int(width_pt / (font_size * 0.92) * 5), 35)
        if current_height([line]) <= max_height_pt:
            pieces.append(line)
        else:
            for start in range(0, len(line), chars_per_piece):
                pieces.append(line[start : start + chars_per_piece])

        for piece in pieces:
            candidate = [*current_lines, piece]
            if current_lines and current_height(candidate) > max_height_pt:
                chunks.append("\n".join(current_lines))
                current_lines = [piece]
            else:
                current_lines = candidate

    if current_lines:
        chunks.append("\n".join(current_lines))
    return chunks


def split_text_into_sentences(text: str) -> list[str]:
    parts: list[str] = []
    current = ""
    for char in text:
        current += char
        if char in "。！？!?；;":
            parts.append(current.strip())
            current = ""
    if current.strip():
        parts.append(current.strip())
    return parts


def split_block_by_capacity(block: dict, width_pt: float, max_height_pt: float, base_size: int = 18) -> list[dict]:
    if estimate_blocks_height([block], width_pt, base_size) <= max_height_pt:
        return [block]
    font_size = block_font_size(block, base_size)
    return [
        make_text_block_like(block, chunk)
        for chunk in split_text_by_height(block_text(block), width_pt, max_height_pt, font_size)
    ]


def split_segments_by_capacity(
    segments: list[list[dict]],
    width_pt: float,
    max_height_pt: float,
    base_size: int = 18,
) -> list[list[dict]]:
    output: list[list[dict]] = []
    for segment in segments:
        current: list[dict] = []
        for block in segment:
            for chunk in split_block_by_capacity(block, width_pt, max_height_pt, base_size):
                candidate = [*current, chunk]
                if current and estimate_blocks_height(candidate, width_pt, base_size) > max_height_pt:
                    output.append(current)
                    current = [chunk]
                else:
                    current = candidate
        if current:
            output.append(current)
    return output


def merge_short_segments(
    segments: list[list[dict]],
    width_pt: float,
    max_height_pt: float,
    base_size: int = 15,
    min_chars: int = 90,
) -> list[list[dict]]:
    merged: list[list[dict]] = []
    for segment in segments:
        segment_chars = blocks_text_length(segment)
        if (
            merged
            and segment_chars < min_chars
            and estimate_blocks_height([*merged[-1], *segment], width_pt, base_size) <= max_height_pt * 1.25
        ):
            merged[-1].extend(segment)
        else:
            merged.append(segment)
    if len(merged) < 2:
        return merged

    output: list[list[dict]] = []
    index = 0
    while index < len(merged):
        segment = merged[index]
        segment_text = "\n".join(block_text(block) for block in segment)
        is_placeholder_only = "图表占位" in segment_text and len(segment_text.strip()) <= 24
        if (
            index + 1 < len(merged)
            and blocks_text_length(segment) < min_chars
            and estimate_blocks_height([*segment, *merged[index + 1]], width_pt, base_size) <= max_height_pt * 1.25
        ):
            output.append([*segment, *merged[index + 1]])
            index += 2
        elif is_placeholder_only and output:
            output[-1].extend(segment)
            index += 1
        else:
            output.append(segment)
            index += 1
    return output


def merge_short_slide_specs(
    specs: list[tuple[list[dict], list[dict]]],
    width_pt: float,
    max_height_pt: float,
    base_size: int = 15,
    min_chars: int = 90,
) -> list[tuple[list[dict], list[dict]]]:
    merged: list[tuple[list[dict], list[dict]]] = []
    for segment, images in specs:
        segment_text = "\n".join(block_text(block) for block in segment).strip()
        segment_chars = len(segment_text)
        is_image_only = not segment_text and bool(images)
        if is_image_only:
            merged.append((segment, images))
            continue
        is_short = segment_chars < min_chars
        is_placeholder_only = "图表占位" in segment_text and segment_chars <= 30
        if merged and (is_short or is_placeholder_only):
            previous_segment, previous_images = merged[-1]
            candidate = [*previous_segment, *segment]
            capacity_multiplier = 1.7 if is_short else 1.25
            max_merged_chars = 900 if is_short else 780
            if blocks_text_length(candidate) <= max_merged_chars and estimate_blocks_height(candidate, width_pt, base_size) <= max_height_pt * capacity_multiplier:
                merged[-1] = (candidate, [*previous_images, *images])
                continue
        merged.append((segment, images))
    return merged


def split_block_for_rebalance(block: dict, max_prefix_chars: int) -> tuple[dict | None, dict | None]:
    text = block_text(block)
    if len(text) <= max_prefix_chars:
        return block, None

    lines = text.splitlines()
    prefix_lines: list[str] = []
    suffix_lines: list[str] = []
    current_chars = 0
    for index, line in enumerate(lines):
        line_len = len(line)
        if prefix_lines and current_chars + line_len > max_prefix_chars:
            suffix_lines = lines[index:]
            break
        if not prefix_lines and line_len > max_prefix_chars:
            split_at = max_prefix_chars
            prefix_lines = [line[:split_at]]
            suffix_lines = [line[split_at:], *lines[index + 1 :]]
            break
        prefix_lines.append(line)
        current_chars += line_len

    if not suffix_lines and len("\n".join(prefix_lines)) == len(text):
        return block, None

    prefix_text = "\n".join(line for line in prefix_lines if line).strip()
    suffix_text = "\n".join(line for line in suffix_lines if line).strip()
    prefix = make_text_block_like(block, prefix_text) if prefix_text else None
    suffix = make_text_block_like(block, suffix_text) if suffix_text else None
    return prefix, suffix


def rebalance_short_slide_specs(
    specs: list[tuple[list[dict], list[dict]]],
    width_pt: float,
    max_height_pt: float,
    base_size: int = 15,
    min_chars: int = 380,
    target_chars: int = 650,
    max_chars: int = 780,
) -> list[tuple[list[dict], list[dict]]]:
    balanced = [(list(segment), list(images)) for segment, images in specs]
    index = 0
    while index + 1 < len(balanced):
        segment, images = balanced[index]
        next_segment, next_images = balanced[index + 1]
        current_text = "\n".join(block_text(block) for block in segment).strip()
        next_text = "\n".join(block_text(block) for block in next_segment).strip()
        current_tail = current_text[-1:] if current_text else ""
        next_head = next_text[:1]
        current_incomplete = current_tail not in "。！？!?：:；;）)】》”’…"
        next_continues = next_head and next_head not in "【一二三四五六七八九十①②③④⑤⑥⑦⑧⑨⑩⑪⑫⑬任务"
        should_rebalance = current_incomplete and next_continues and blocks_text_length(segment) < min_chars

        if images or not next_segment or not should_rebalance:
            index += 1
            continue

        moved_any = False
        while next_segment and blocks_text_length(segment) < target_chars:
            remaining_capacity = max_chars - blocks_text_length(segment)
            if remaining_capacity <= 80:
                break
            prefix, suffix = split_block_for_rebalance(next_segment[0], remaining_capacity)
            if prefix is None:
                break
            candidate = [*segment, prefix]
            if estimate_blocks_height(candidate, width_pt, base_size) > max_height_pt * 1.25:
                break
            segment = candidate
            moved_any = True
            if suffix is None:
                next_segment = next_segment[1:]
            else:
                next_segment = [suffix, *next_segment[1:]]
                break

        if moved_any:
            balanced[index] = (segment, images)
            if next_segment:
                balanced[index + 1] = (next_segment, next_images)
            else:
                if next_images:
                    balanced[index] = (segment, [*images, *next_images])
                del balanced[index + 1]
                continue
        index += 1
    return balanced


def split_long_segments(segments: list[list[dict]], max_chars: int = 850) -> list[list[dict]]:
    output: list[list[dict]] = []
    for segment in segments:
        current: list[dict] = []
        current_chars = 0
        for block in split_oversized_blocks(segment, max_chars):
            text_len = len(block_text(block))
            if current and current_chars + text_len > max_chars:
                output.append(current)
                current = []
                current_chars = 0
            current.append(block)
            current_chars += text_len
        if current:
            output.append(current)
    return output


def split_text_by_size(text: str, max_chars: int) -> list[str]:
    chunks: list[str] = []
    current = ""
    pieces: list[str] = []
    for line in text.splitlines():
        line = line.strip()
        if not line:
            continue
        start = 0
        while start < len(line):
            pieces.append(line[start : start + max_chars])
            start += max_chars

    for piece in pieces:
        separator = "\n" if current else ""
        if current and len(current) + len(separator) + len(piece) > max_chars:
            chunks.append(current)
            current = piece
        else:
            current = f"{current}{separator}{piece}" if current else piece
    if current:
        chunks.append(current)
    return chunks


def split_oversized_blocks(blocks: list[dict], max_chars: int) -> list[dict]:
    output: list[dict] = []
    for block in blocks:
        text = block_text(block)
        if len(text) <= max_chars:
            output.append(block)
            continue
        for chunk in split_text_by_size(text, max_chars):
            output.append(make_text_block_like(block, chunk))
    return output


def first_span(block: dict) -> dict | None:
    if "_span" in block:
        return block["_span"]

    for line in block.get("lines", []):
        for span in line.get("spans", []):
            if span.get("text", "").strip():
                return span
    return None


def choose_font_size(chars: int) -> int:
    if chars <= 180:
        return 22
    if chars <= 320:
        return 18
    return 15


def is_answer_animation_block(block: dict) -> bool:
    return block.get("_role") in {"answer_label", "answer"}


def segment_has_animated_answers(segment: list[dict]) -> bool:
    return any(is_answer_animation_block(block) for block in segment)


def grouped_blocks_for_answer_animation(segment: list[dict]) -> list[tuple[list[dict], bool]]:
    groups: list[tuple[list[dict], bool]] = []
    index = 0
    while index < len(segment):
        block = segment[index]
        role = block.get("_role")
        if role == "answer_label":
            group = [block]
            index += 1
            while index < len(segment) and segment[index].get("_role") == "answer":
                group.append(segment[index])
                index += 1
            groups.append((group, True))
            continue
        if role == "answer":
            groups.append(([block], True))
            index += 1
            continue

        group = [block]
        index += 1
        while (
            index < len(segment)
            and not is_answer_animation_block(segment[index])
            and should_merge_text_blocks(block_text(group[-1]), block_text(segment[index]))
        ):
            group.append(segment[index])
            index += 1
        groups.append((group, False))
    return groups


def format_text_run(run, block: dict, base_size: int) -> None:
    span = first_span(block) or {}
    text = block_text(block)
    source_size = span.get("size", 10)
    is_heading = source_size >= 13 or span_is_bold(span)
    font = run.font
    font.name = clean_font_name(span.get("font", "Arial"))
    font.size = Pt(block_font_size(block, base_size))
    font.bold = block.get("_bold", is_heading)
    font.italic = span_is_italic(span)
    font.color.rgb = int_to_rgb(span.get("color", 0))


def add_text_group_to_box(frame, blocks: list[dict], base_size: int) -> None:
    first_paragraph = True
    previous_paragraph = None
    for block in blocks:
        text = block_text(block)
        if not text:
            continue
        if previous_paragraph is not None and should_merge_text_blocks(previous_paragraph.text.strip(), text):
            paragraph = previous_paragraph
        else:
            paragraph = frame.paragraphs[0] if first_paragraph else frame.add_paragraph()
            first_paragraph = False
            paragraph.alignment = PP_ALIGN.LEFT
            paragraph.space_before = Pt(0)
            paragraph.space_after = Pt(5)
            paragraph.line_spacing = 1.0
            previous_paragraph = paragraph

        run = paragraph.add_run()
        run.text = text
        format_text_run(run, block, base_size)


def add_fly_in_animations(slide, shape_ids: list[int]) -> None:
    if not shape_ids:
        return

    existing_timing = slide._element.find(qn("p:timing"))
    if existing_timing is not None:
        slide._element.remove(existing_timing)

    next_timing_id = 3
    effect_nodes: list[str] = []
    build_nodes: list[str] = []
    for shape_id in shape_ids:
        wrapper_id = next_timing_id
        effect_id = next_timing_id + 1
        visibility_id = next_timing_id + 2
        fly_id = next_timing_id + 3
        next_timing_id += 4
        effect_nodes.append(
            f"""
            <p:par>
              <p:cTn id="{wrapper_id}" fill="hold">
                <p:stCondLst><p:cond delay="0"/></p:stCondLst>
                <p:childTnLst>
                  <p:par>
                    <p:cTn id="{effect_id}" presetID="2" presetClass="entr" presetSubtype="4" fill="hold" grpId="0" nodeType="clickEffect">
                      <p:stCondLst><p:cond delay="0"/></p:stCondLst>
                      <p:childTnLst>
                        <p:set>
                          <p:cBhvr>
                            <p:cTn id="{visibility_id}" dur="1" fill="hold"/>
                            <p:tgtEl><p:spTgt spid="{shape_id}"/></p:tgtEl>
                            <p:attrNameLst><p:attrName>style.visibility</p:attrName></p:attrNameLst>
                          </p:cBhvr>
                          <p:to><p:strVal val="visible"/></p:to>
                        </p:set>
                        <p:animEffect transition="in" filter="fly" prLst="dir=l">
                          <p:cBhvr>
                            <p:cTn id="{fly_id}" dur="{ANSWER_ANIMATION_DURATION_MS}"/>
                            <p:tgtEl><p:spTgt spid="{shape_id}"/></p:tgtEl>
                          </p:cBhvr>
                        </p:animEffect>
                      </p:childTnLst>
                    </p:cTn>
                  </p:par>
                </p:childTnLst>
              </p:cTn>
            </p:par>
            """
        )
        build_nodes.append(f'<p:bldP spid="{shape_id}" grpId="0"/>')

    timing = parse_xml(
        f"""
        <p:timing {nsdecls("p")}>
          <p:tnLst>
            <p:par>
              <p:cTn id="1" dur="indefinite" restart="never" nodeType="tmRoot">
                <p:childTnLst>
                  <p:seq concurrent="1" nextAc="seek">
                    <p:cTn id="2" dur="indefinite" nodeType="mainSeq">
                      <p:childTnLst>
                        {"".join(effect_nodes)}
                      </p:childTnLst>
                    </p:cTn>
                  </p:seq>
                </p:childTnLst>
              </p:cTn>
            </p:par>
          </p:tnLst>
          <p:bldLst>{"".join(build_nodes)}</p:bldLst>
        </p:timing>
        """
    )
    ext_lst = slide._element.find(qn("p:extLst"))
    if ext_lst is None:
        slide._element.append(timing)
    else:
        slide._element.insert(slide._element.index(ext_lst), timing)


def add_animated_answer_text_slide(
    slide,
    segment: list[dict],
    slide_width_pt: float,
    slide_height_pt: float,
    image_area_height: float,
    margin_x: float,
    margin_y: float,
) -> None:
    base_size = choose_font_size(blocks_text_length(segment))
    content_width = slide_width_pt - margin_x * 2
    text_width = content_width
    max_text_height = slide_height_pt - margin_y * 2 - image_area_height
    current_top = margin_y
    answer_shape_ids: list[int] = []

    for blocks, animate in grouped_blocks_for_answer_animation(segment):
        group_height = estimate_blocks_height(
            blocks,
            text_width - TEXT_INSET_X_PT * 2,
            base_size,
        )
        box_height = max(18, group_height + TEXT_INSET_Y_PT * 2)
        if current_top + box_height > margin_y + max_text_height:
            box_height = max(18, margin_y + max_text_height - current_top)
        box = slide.shapes.add_textbox(
            pt_to_emu(margin_x),
            pt_to_emu(current_top),
            pt_to_emu(text_width),
            pt_to_emu(box_height),
        )
        frame = box.text_frame
        frame.clear()
        frame.margin_left = pt_to_emu(TEXT_INSET_X_PT)
        frame.margin_right = pt_to_emu(TEXT_INSET_X_PT)
        frame.margin_top = pt_to_emu(TEXT_INSET_Y_PT)
        frame.margin_bottom = pt_to_emu(TEXT_INSET_Y_PT)
        frame.word_wrap = True
        frame.auto_size = MSO_AUTO_SIZE.TEXT_TO_FIT_SHAPE
        frame.vertical_anchor = MSO_ANCHOR.TOP
        add_text_group_to_box(frame, blocks, base_size)
        if animate:
            answer_shape_ids.append(box.shape_id)
        current_top += box_height

    add_fly_in_animations(slide, answer_shape_ids)


def add_clean_text_slide(slide, segment: list[dict], slide_width_pt: float, slide_height_pt: float, images: list[dict] | None = None) -> None:
    images = images or []
    has_text_content = any(block_text(block).strip() for block in segment)
    margin_x = CONTENT_MARGIN_X_PT
    margin_y = CONTENT_MARGIN_Y_PT
    image_area_height = 0
    if images and has_text_content:
        image_area_height = min(IMAGE_AREA_HEIGHT_PT, slide_height_pt * 0.36)

    if has_text_content:
        if segment_has_animated_answers(segment):
            add_animated_answer_text_slide(slide, segment, slide_width_pt, slide_height_pt, image_area_height, margin_x, margin_y)
        else:
            box = slide.shapes.add_textbox(
                pt_to_emu(margin_x),
                pt_to_emu(margin_y),
                pt_to_emu(slide_width_pt - margin_x * 2),
                pt_to_emu(slide_height_pt - margin_y * 2 - image_area_height),
            )
            frame = box.text_frame
            frame.clear()
            frame.margin_left = pt_to_emu(TEXT_INSET_X_PT)
            frame.margin_right = pt_to_emu(TEXT_INSET_X_PT)
            frame.margin_top = pt_to_emu(TEXT_INSET_Y_PT)
            frame.margin_bottom = pt_to_emu(TEXT_INSET_Y_PT)
            frame.word_wrap = True
            frame.auto_size = MSO_AUTO_SIZE.TEXT_TO_FIT_SHAPE
            frame.vertical_anchor = MSO_ANCHOR.TOP
            add_text_group_to_box(frame, segment, choose_font_size(blocks_text_length(segment)))

    add_images_to_slide(
        slide,
        images,
        slide_width_pt,
        slide_height_pt,
        margin_x,
        margin_y,
        center_vertical=not has_text_content,
    )


def image_display_size_pt(image_block: dict) -> tuple[float, float]:
    x0, y0, x1, y1 = object_bbox(image_block)
    bbox_width = max(x1 - x0, 1)
    bbox_height = max(y1 - y0, 1)
    return bbox_width, bbox_height


def add_images_to_slide(
    slide,
    image_blocks: list[dict],
    slide_width_pt: float,
    slide_height_pt: float,
    margin_x: float,
    margin_y: float,
    center_vertical: bool = False,
) -> None:
    if not image_blocks:
        return

    gap = 12
    max_total_width = slide_width_pt - margin_x * 2
    max_height = slide_height_pt - margin_y * 2 if center_vertical else min(175, slide_height_pt * 0.32)
    slots = len(image_blocks)
    slot_width = (max_total_width - gap * (slots - 1)) / slots
    rendered: list[tuple[dict, float, float]] = []

    for image_block in image_blocks:
        image_bytes = image_block.get("image")
        if not image_bytes:
            continue
        width, height = image_display_size_pt(image_block)
        max_upscale = 3.0 if center_vertical else 1.0
        scale = min(slot_width / width, max_height / height, max_upscale)
        rendered.append((image_block, width * scale, height * scale))

    if not rendered:
        return

    total_width = sum(width for _, width, _ in rendered) + gap * (len(rendered) - 1)
    row_height = max(height for _, _, height in rendered)
    left = (slide_width_pt - total_width) / 2
    if center_vertical:
        top = (slide_height_pt - row_height) / 2
    else:
        top = slide_height_pt - margin_y - row_height
    for image_block, width, height in rendered:
        slide.shapes.add_picture(
            io.BytesIO(image_block["image"]),
            pt_to_emu(left),
            pt_to_emu(top + (row_height - height) / 2 if center_vertical else top + (max_height - height) / 2),
            width=pt_to_emu(width),
            height=pt_to_emu(height),
        )
        left += width + gap


def add_positioned_image_block(
    slide,
    image_block: dict,
    scale: float,
    offset_x: float = 0,
    offset_y: float = 0,
    margin_left: float = 0,
    margin_top: float = 0,
) -> None:
    image_bytes = image_block.get("image")
    if not image_bytes:
        return
    x0, y0, x1, y1 = object_bbox(image_block)
    left = pt_to_emu((x0 - offset_x) * scale) + pt_to_emu(margin_left)
    top = pt_to_emu((y0 - offset_y) * scale) + pt_to_emu(margin_top)
    width = max(pt_to_emu((x1 - x0) * scale), pt_to_emu(6))
    height = max(pt_to_emu((y1 - y0) * scale), pt_to_emu(6))
    slide.shapes.add_picture(io.BytesIO(image_bytes), left, top, width=width, height=height)


def detect_answer_start_page(doc: fitz.Document, profile: ConversionProfile) -> int | None:
    for index, page in enumerate(doc, 1):
        if index < profile.answer_min_scan_page:
            continue
        text = page.get_text()
        if any(marker in text for marker in profile.answer_markers) or text.count(profile.answer_keyword) >= profile.answer_keyword_count:
            return index
    return profile.answer_fallback_page


def convert_pdf_to_pptx(
    pdf_path: Path,
    pptx_path: Path,
    skip_from_page: int | None = 29,
    layout: str = "clean",
    profile: str | ConversionProfile = "workbook",
    answer_mode: str = "skip",
) -> ConversionSummary:
    if layout not in {"clean", "positioned"}:
        raise ValueError(f"Unsupported layout: {layout}")
    if answer_mode not in {"skip", "inline"}:
        raise ValueError(f"Unsupported answer mode: {answer_mode}")
    if answer_mode == "inline" and layout != "clean":
        raise ValueError("Inline answer mode only supports clean layout")

    conversion_profile = resolve_profile(profile)
    doc = fitz.open(pdf_path)
    try:
        if doc.page_count == 0:
            raise ValueError(f"PDF has no pages: {pdf_path}")

        source_pages = doc.page_count
        slide_width_pt = 13.333333 * PDF_POINTS_PER_INCH
        slide_height_pt = 7.5 * PDF_POINTS_PER_INCH

        presentation = Presentation()
        presentation.slide_width = Inches(13.333333)
        presentation.slide_height = Inches(7.5)
        blank_layout = presentation.slide_layouts[6]

        if skip_from_page == 0:
            skip_from_page = detect_answer_start_page(doc, conversion_profile)

        if answer_mode == "inline":
            answer_start_page = skip_from_page or detect_answer_start_page(doc, conversion_profile)
            if answer_start_page is None:
                raise ValueError("Could not detect answer pages for inline answer mode")
            if answer_start_page <= 1 or answer_start_page > doc.page_count:
                raise ValueError(f"Invalid answer start page: {answer_start_page}")
            slide_specs, topic_count, matched_answers, unmatched_topics, unused_answers = build_inline_answer_slide_specs(
                doc,
                answer_start_page,
                slide_width_pt,
                slide_height_pt,
            )
            for segment, images in slide_specs:
                slide = presentation.slides.add_slide(blank_layout)
                add_clean_text_slide(slide, segment, slide_width_pt, slide_height_pt, images)
            pptx_path.parent.mkdir(parents=True, exist_ok=True)
            presentation.save(pptx_path)
            return ConversionSummary(
                pdf_path=pdf_path,
                pptx_path=pptx_path,
                source_pages=source_pages,
                processed_pages=answer_start_page - 1,
                skipped_from_page=answer_start_page,
                slides=len(presentation.slides),
                layout=layout,
                answer_mode=answer_mode,
                matched_answers=matched_answers,
                unmatched_topics=unmatched_topics,
                unused_answers=unused_answers,
            )

        processed_pages = 0
        raw_slide_specs: list[tuple[list[dict], list[dict], fitz.Rect]] = []

        for page_number, page in enumerate(doc, 1):
            if skip_from_page is not None and page_number >= skip_from_page:
                break
            processed_pages += 1
            page_dict = page.get_text("dict", flags=fitz.TEXT_PRESERVE_LIGATURES | fitz.TEXT_PRESERVE_WHITESPACE)
            image_page_dict = page.get_text("dict")
            text_blocks = [block for block in page_dict.get("blocks", []) if block.get("type") == 0 and has_text(block)]
            text_blocks = replace_table_text_with_placeholders(text_blocks, conversion_profile)
            text_blocks = merge_adjacent_text_blocks(text_blocks)
            image_blocks = [block for block in image_page_dict.get("blocks", []) if block.get("type") == 1]
            has_images = bool(image_blocks)
            content_width_pt = slide_width_pt - CONTENT_MARGIN_X_PT * 2 - TEXT_INSET_X_PT * 2
            visible_height_pt = (
                slide_height_pt
                - CONTENT_MARGIN_Y_PT * 2
                - TEXT_INSET_Y_PT * 2
                - (IMAGE_AREA_HEIGHT_PT if has_images and text_blocks else 0)
            )
            content_height_pt = visible_height_pt * CAPACITY_SAFETY
            segments = split_segments_by_capacity(
                split_blocks_by_vertical_gaps(text_blocks),
                width_pt=content_width_pt,
                max_height_pt=content_height_pt,
                base_size=15,
            )
            segments = merge_short_segments(segments, content_width_pt, content_height_pt, base_size=15)

            if segments:
                image_assignments = assign_images_to_segments(image_blocks, segments, fitz.Rect(page.rect))
                for segment_index, segment in enumerate(segments):
                    raw_slide_specs.append((segment, image_assignments[segment_index], fitz.Rect(page.rect)))
            elif image_blocks:
                raw_slide_specs.append(([], image_blocks, fitz.Rect(page.rect)))

        if layout == "clean":
            slide_specs = [(segment, images) for segment, images, _ in raw_slide_specs]
            slide_specs = merge_short_slide_specs(
                slide_specs,
                width_pt=slide_width_pt - CONTENT_MARGIN_X_PT * 2 - TEXT_INSET_X_PT * 2,
                max_height_pt=(slide_height_pt - CONTENT_MARGIN_Y_PT * 2 - TEXT_INSET_Y_PT * 2) * CAPACITY_SAFETY,
                base_size=15,
            )
            slide_specs = rebalance_short_slide_specs(
                slide_specs,
                width_pt=slide_width_pt - CONTENT_MARGIN_X_PT * 2 - TEXT_INSET_X_PT * 2,
                max_height_pt=(slide_height_pt - CONTENT_MARGIN_Y_PT * 2 - TEXT_INSET_Y_PT * 2) * CAPACITY_SAFETY,
                base_size=15,
                min_chars=380,
                target_chars=520,
                max_chars=760,
            )
            slide_specs = merge_short_slide_specs(
                slide_specs,
                width_pt=slide_width_pt - CONTENT_MARGIN_X_PT * 2 - TEXT_INSET_X_PT * 2,
                max_height_pt=(slide_height_pt - CONTENT_MARGIN_Y_PT * 2 - TEXT_INSET_Y_PT * 2) * CAPACITY_SAFETY,
                base_size=15,
            )

            for segment, images in slide_specs:
                if not segment and not images:
                    continue
                slide = presentation.slides.add_slide(blank_layout)
                add_clean_text_slide(
                    slide,
                    segment,
                    slide_width_pt,
                    slide_height_pt,
                    images,
                )
        else:
            for segment, images, page_rect in raw_slide_specs:
                if not segment and not images:
                    continue
                slide = presentation.slides.add_slide(blank_layout)
                x0, y0, x1, y1 = spec_bbox(segment, images, page_rect)
                segment_width = max(x1 - x0, 1)
                segment_height = max(y1 - y0, 1)
                slide_margin_x = 34
                slide_margin_y = 28
                scale = min(
                    (slide_width_pt - slide_margin_x * 2) / segment_width,
                    (slide_height_pt - slide_margin_y * 2) / segment_height,
                )
                scale = min(scale, 2.2)
                used_width = segment_width * scale
                used_height = segment_height * scale
                margin_left = max(slide_margin_x, (slide_width_pt - used_width) / 2)
                margin_top = max(slide_margin_y, (slide_height_pt - used_height) / 2)

                for image_block in images:
                    add_positioned_image_block(slide, image_block, scale, x0, y0, margin_left, margin_top)
                for block in segment:
                    add_text_block(slide, block, scale, x0, y0, margin_left, margin_top)

        pptx_path.parent.mkdir(parents=True, exist_ok=True)
        presentation.save(pptx_path)
        return ConversionSummary(
            pdf_path=pdf_path,
            pptx_path=pptx_path,
            source_pages=source_pages,
            processed_pages=processed_pages,
            skipped_from_page=skip_from_page,
            slides=len(presentation.slides),
            layout=layout,
            answer_mode=answer_mode,
        )
    finally:
        doc.close()


def main() -> None:
    parser = argparse.ArgumentParser(description="Convert copyable PDF text into an editable PPTX.")
    parser.add_argument("pdf", nargs="?", default="mypaper.pdf", help="Input PDF path")
    parser.add_argument("pptx", nargs="?", default="mypaper.pptx", help="Output PPTX path")
    parser.add_argument("--skip-from-page", type=int, default=29, help="Skip this PDF page and all following pages. Use 0 to auto-detect, -1 to keep all pages.")
    parser.add_argument("--layout", choices=("clean", "positioned"), default="clean", help="clean avoids overlap; positioned preserves original coordinates.")
    parser.add_argument("--profile", choices=tuple(sorted(CONVERSION_PROFILES)), default="workbook", help="Rule profile for answer-page and chart detection.")
    parser.add_argument("--answer-mode", choices=("skip", "inline"), default="skip", help="skip omits answer pages; inline places matched answers below each topic.")
    args = parser.parse_args()

    skip_from_page = None if args.skip_from_page < 0 else args.skip_from_page
    convert_pdf_to_pptx(Path(args.pdf), Path(args.pptx), skip_from_page=skip_from_page, layout=args.layout, profile=args.profile, answer_mode=args.answer_mode)


if __name__ == "__main__":
    main()
