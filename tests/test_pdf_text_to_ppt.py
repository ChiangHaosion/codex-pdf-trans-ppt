from __future__ import annotations

import io
import re
import zipfile
from pathlib import Path

import fitz
from lxml import etree
from PIL import Image
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE_TYPE

from pdf_text_to_ppt import (
    CONVERSION_PROFILES,
    block_text,
    convert_pdf_to_pptx,
    join_pdf_lines,
    replace_table_text_with_placeholders,
)


def text_block(text: str) -> dict:
    return {
        "bbox": (0, 0, 220, 24),
        "lines": [
            {
                "spans": [
                    {
                        "text": text,
                        "size": 10,
                        "font": "Arial",
                        "color": 0,
                    }
                ]
            }
        ],
    }


def make_png_bytes() -> bytes:
    image = Image.new("RGB", (80, 48), (36, 107, 254))
    stream = io.BytesIO()
    image.save(stream, format="PNG")
    return stream.getvalue()


def make_image_pdf(path: Path) -> None:
    doc = fitz.open()
    page = doc.new_page(width=320, height=220)
    page.insert_image(fitz.Rect(60, 54, 220, 150), stream=make_png_bytes())
    doc.save(path)
    doc.close()


def count_pictures(pptx_path: Path) -> int:
    presentation = Presentation(pptx_path)
    return sum(
        1
        for slide in presentation.slides
        for shape in slide.shapes
        if shape.shape_type == MSO_SHAPE_TYPE.PICTURE
    )


def pptx_text(pptx_path: Path) -> str:
    return "\n".join(pptx_slide_texts(pptx_path))


def pptx_slide_texts(pptx_path: Path) -> list[str]:
    presentation = Presentation(pptx_path)
    return [
        "\n".join(
            shape.text
            for shape in slide.shapes
            if hasattr(shape, "text") and shape.text
        )
        for slide in presentation.slides
    ]


def pptx_runs(pptx_path: Path):
    presentation = Presentation(pptx_path)
    for slide in presentation.slides:
        for shape in slide.shapes:
            if not hasattr(shape, "text_frame"):
                continue
            for paragraph in shape.text_frame.paragraphs:
                for run in paragraph.runs:
                    yield run


def text_between(text: str, start_marker: str, end_marker: str) -> str:
    start = text.index(start_marker)
    end = text.index(end_marker, start + len(start_marker))
    return text[start:end]


def assert_no_inline_answer_orphans(slide_texts: list[str]) -> None:
    labels = ("例题参考答案", "巩固训练参考答案", "参考答案")
    for slide_text in slide_texts:
        compact = "".join(slide_text.split())
        assert not re.fullmatch(r"专题[一二三四五六七八九十]+--.+（\d+）", compact)
        assert len(compact) >= 90
        for label in labels:
            assert compact != label
            assert not slide_text.rstrip().endswith(label)

    for slide_text in slide_texts:
        for label in ("例题参考答案", "巩固训练参考答案"):
            start = slide_text.find(label)
            while start >= 0:
                after = slide_text[start + len(label) :].strip()
                next_labels = [after.find(other) for other in labels if after.find(other) >= 0]
                if next_labels:
                    after = after[: min(next_labels)].strip()
                assert len("".join(after.split())) >= 8
                start = slide_text.find(label, start + 1)


def assert_inline_answers_are_red(pptx_path: Path) -> None:
    expected = RGBColor(0xC6, 0x28, 0x28)
    checked = set()
    for run in pptx_runs(pptx_path):
        if "例题参考答案" in run.text:
            assert run.font.color.rgb == expected
            checked.add("example_label")
        if "巩固训练参考答案" in run.text:
            assert run.font.color.rgb == expected
            checked.add("training_label")
        if "燕子捕食害虫数量之多" in run.text:
            assert run.font.color.rgb == expected
            checked.add("answer_body")

    assert checked == {"example_label", "training_label", "answer_body"}


def assert_inline_answers_have_fly_in_animation(pptx_path: Path) -> None:
    namespaces = {
        "a": "http://schemas.openxmlformats.org/drawingml/2006/main",
        "p": "http://schemas.openxmlformats.org/presentationml/2006/main",
    }
    animated_answer_shapes = 0
    with zipfile.ZipFile(pptx_path) as archive:
        slide_names = sorted(
            name for name in archive.namelist() if re.fullmatch(r"ppt/slides/slide\d+\.xml", name)
        )
        for slide_name in slide_names:
            root = etree.fromstring(archive.read(slide_name))
            target_ids = set(
                root.xpath(
                    './/p:anim[p:cBhvr/p:attrNameLst/p:attrName="ppt_x"]/p:cBhvr/p:tgtEl/p:spTgt/@spid',
                    namespaces=namespaces,
                )
            )
            if not target_ids:
                continue
            assert root.xpath('.//p:cTn/p:stCondLst/p:cond[@delay="indefinite"]', namespaces=namespaces)
            assert set(root.xpath('.//p:anim/p:cBhvr/p:attrNameLst/p:attrName/text()', namespaces=namespaces)) >= {
                "ppt_x",
                "ppt_y",
            }
            assert root.xpath('.//p:bldP[@advAuto="indefinite" and @build="whole" and @animBg="1"]', namespaces=namespaces)
            for shape in root.xpath(".//p:sp", namespaces=namespaces):
                shape_id = shape.xpath("./p:nvSpPr/p:cNvPr/@id", namespaces=namespaces)[0]
                text = "".join(shape.xpath(".//a:t/text()", namespaces=namespaces))
                colors = set(shape.xpath(".//a:srgbClr/@val", namespaces=namespaces))
                is_answer_shape = "参考答案" in text or "C62828" in colors
                if is_answer_shape:
                    assert shape_id in target_ids
                    animated_answer_shapes += 1
                elif text.strip():
                    assert shape_id not in target_ids

    assert animated_answer_shapes >= 31


def assert_long_topic_answers_follow_questions(deck_text: str) -> None:
    topic_text = text_between(deck_text, "专题二--记叙文（1）", "专题二--记叙文（2）")
    assert topic_text.index("本文以“一条怕水的鱼”为题，有何妙处？") < topic_text.index("例题参考答案")
    assert topic_text.index("例题参考答案") < topic_text.index("【巩固训练】")
    assert topic_text.index("作者以“美丽的萝卜花”为题有何用意?") < topic_text.index("巩固训练参考答案")


def test_join_pdf_lines_keeps_structural_breaks() -> None:
    text = join_pdf_lines(["阅读答题技巧小纸条", "一、阅读材料", "这是", "正文"])

    assert text == "阅读答题技巧小纸条\n一、阅读材料\n这是正文"


def test_table_placeholder_uses_profile_rules() -> None:
    profile = CONVERSION_PROFILES["workbook"]

    output = replace_table_text_with_placeholders(
        [text_block("用途 所占比例 10% 20% 30%"), text_block("经常使用 1234")],
        profile,
    )

    assert block_text(output[0]) == profile.table_placeholder
    assert len(output) == 1


def test_image_only_pdf_generates_clean_slide_with_picture(tmp_path: Path) -> None:
    pdf_path = tmp_path / "image_only.pdf"
    pptx_path = tmp_path / "image_only.pptx"
    make_image_pdf(pdf_path)

    summary = convert_pdf_to_pptx(pdf_path, pptx_path, skip_from_page=None, layout="clean", profile="generic")

    assert summary.source_pages == 1
    assert summary.processed_pages == 1
    assert summary.slides == 1
    assert count_pictures(pptx_path) == 1


def test_positioned_layout_keeps_image_blocks(tmp_path: Path) -> None:
    pdf_path = tmp_path / "image_positioned.pdf"
    pptx_path = tmp_path / "image_positioned.pptx"
    make_image_pdf(pdf_path)

    summary = convert_pdf_to_pptx(pdf_path, pptx_path, skip_from_page=None, layout="positioned", profile="generic")

    assert summary.slides == 1
    assert count_pictures(pptx_path) == 1


def test_sample_workbook_conversion_regression(tmp_path: Path) -> None:
    sample_pdf = Path(__file__).resolve().parents[1] / "mypaper.pdf"
    pptx_path = tmp_path / "sample.pptx"

    summary = convert_pdf_to_pptx(sample_pdf, pptx_path, skip_from_page=29, layout="clean", profile="workbook")

    assert summary.source_pages == 35
    assert summary.processed_pages == 28
    assert summary.slides == 60
    assert pptx_path.exists()


def test_sample_workbook_inline_answers_regression(tmp_path: Path) -> None:
    sample_pdf = Path(__file__).resolve().parents[1] / "mypaper.pdf"
    pptx_path = tmp_path / "sample_answers.pptx"

    summary = convert_pdf_to_pptx(
        sample_pdf,
        pptx_path,
        skip_from_page=29,
        layout="clean",
        profile="workbook",
        answer_mode="inline",
    )
    text = pptx_text(pptx_path)
    slide_texts = pptx_slide_texts(pptx_path)

    assert summary.source_pages == 35
    assert summary.processed_pages == 28
    assert summary.slides == 78
    assert summary.answer_mode == "inline"
    assert summary.matched_answers == 31
    assert summary.unmatched_topics == 0
    assert "例题参考答案" in text
    assert "巩固训练参考答案" in text
    assert "燕子捕食害虫数量之多" in text
    assert_no_inline_answer_orphans(slide_texts)
    assert_inline_answers_are_red(pptx_path)
    assert_inline_answers_have_fly_in_animation(pptx_path)
    assert_long_topic_answers_follow_questions(text)
